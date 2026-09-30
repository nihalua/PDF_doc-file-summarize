"""Read an uploaded file (PDF, DOCX, DOC, TXT) into plain text, locally.

Nothing here touches the network: PDFs are read with pypdf, Word files with
python-docx, and legacy .doc files are converted by a local LibreOffice.
"""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".doc", ".txt", ".md")

# Below this many extracted characters per page, treat the PDF as scanned.
SCANNED_CHARS_PER_PAGE = 30


class ExtractionError(Exception):
    """The file couldn't be read."""


@dataclass
class LoadedDocument:
    filename: str
    # One entry per page for PDFs; a single entry for other formats.
    pages: list[str]
    page_count: int | None = None  # set for PDFs only

    @property
    def text(self) -> str:
        return "\n\n".join(self.pages)


def load_document(filename: str, data: bytes) -> LoadedDocument:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _load_pdf(filename, data)
    if ext == ".docx":
        return LoadedDocument(filename, [_docx_to_text(data)])
    if ext == ".doc":
        return LoadedDocument(filename, [_doc_to_text(data)])
    if ext in (".txt", ".md"):
        return LoadedDocument(filename, [_decode_text(data)])
    raise ExtractionError(
        f"Unsupported file type '{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
    )


def _load_pdf(filename: str, data: bytes) -> LoadedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("This PDF is password-protected.")
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except ExtractionError:
        raise
    except Exception as e:  # pypdf raises a variety of errors on broken files
        raise ExtractionError(f"Could not read PDF: {e}") from e

    if not pages:
        raise ExtractionError("The PDF has no pages.")
    if sum(len(p) for p in pages) < SCANNED_CHARS_PER_PAGE * len(pages):
        raise ExtractionError(
            "This PDF seems to be scanned images without a text layer, so there is no text "
            "to summarize. Run it through OCR first (for example \"Recognize text\" in "
            "Adobe Acrobat or the free OCRmyPDF) and upload the result."
        )
    return LoadedDocument(filename, pages, page_count=len(pages))


def _docx_to_text(data: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as e:
        raise ExtractionError(f"Could not read DOCX: {e}") from e

    parts: list[str] = []
    # iter_inner_content keeps paragraphs and tables in document order.
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "").lower()
            if style.startswith(("heading", "title", "başlık", "baslik")):
                text = f"## {text}"
            elif style.startswith("list"):
                text = f"- {text}"
            parts.append(text)
        elif isinstance(block, Table):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
            if rows:
                parts.append("\n".join(rows))

    text = "\n\n".join(parts).strip()
    if not text:
        raise ExtractionError("The DOCX file contains no readable text.")
    return text


def _doc_to_text(data: bytes) -> str:
    """Legacy .doc: convert to .docx with LibreOffice, then read that."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ExtractionError(
            "Reading legacy .doc files requires LibreOffice (soffice) to be installed. "
            "Alternatively, save the file as .docx or .pdf and upload that."
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "input.doc"
        src.write_bytes(data)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)],
                check=True,
                capture_output=True,
                timeout=180,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            raise ExtractionError(f"LibreOffice could not convert the .doc file: {e}") from e
        out = Path(tmp) / "input.docx"
        if not out.exists():
            raise ExtractionError(
                "LibreOffice could not convert the .doc file (is the libreoffice-writer "
                "package installed?). You can also save it as .docx or .pdf and upload that."
            )
        return _docx_to_text(out.read_bytes())


def _decode_text(data: bytes) -> str:
    # Turkish files are often cp1254 (Windows) or iso-8859-9 rather than UTF-8.
    for enc in ("utf-8-sig", "cp1254", "iso-8859-9"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    if not text.strip():
        raise ExtractionError("The file is empty.")
    return text
