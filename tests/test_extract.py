import io
import shutil
import subprocess

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from docsum.extract import ExtractionError, load_document


def make_pdf(pages: int, text: str = "Hello page") -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i in range(pages):
        c.drawString(72, 720, f"{text} {i + 1} " + "lorem ipsum dolor sit amet " * 3)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx() -> bytes:
    d = Document()
    d.add_heading("Sözleşme Başlığı", level=1)
    d.add_paragraph("Bu sözleşme iki taraf arasında yapılmıştır.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tutar"
    table.cell(0, 1).text = "10.000 TL"
    table.cell(1, 0).text = "Tarih"
    table.cell(1, 1).text = "1 Ocak 2027"
    d.add_paragraph("Son paragraf.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_small_pdf_is_sent_natively():
    doc = load_document("a.pdf", make_pdf(3))
    assert doc.is_pdf and len(doc.pdf_batches) == 1 and doc.page_count == 3


def test_large_text_pdf_is_extracted_to_text():
    doc = load_document("big.pdf", make_pdf(120))
    assert not doc.is_pdf
    assert "[Page 120]" in doc.text and "Hello page 1 " in doc.text


def test_large_scanned_pdf_is_split(monkeypatch):
    from pypdf import PageObject
    monkeypatch.setattr(PageObject, "extract_text", lambda self, *a, **k: "")
    doc = load_document("scan.pdf", make_pdf(250))
    assert doc.is_pdf and len(doc.pdf_batches) == 3


def test_docx_keeps_order_headings_and_tables():
    text = load_document("s.docx", make_docx()).text
    assert text.index("## Sözleşme Başlığı") < text.index("Tutar | 10.000 TL") < text.index("Son paragraf")


def test_turkish_cp1254_text():
    text = load_document("t.txt", "Şirket ğüşıöç".encode("cp1254")).text
    assert text == "Şirket ğüşıöç"


def test_unsupported_and_broken_files():
    with pytest.raises(ExtractionError):
        load_document("x.xlsx", b"123")
    with pytest.raises(ExtractionError):
        load_document("x.pdf", b"not a pdf")
    with pytest.raises(ExtractionError):
        load_document("x.docx", b"not a docx")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_legacy_doc(tmp_path):
    src = tmp_path / "s.docx"
    src.write_bytes(make_docx())
    subprocess.run(["soffice", "--headless", "--convert-to", "doc:MS Word 97", "--outdir", str(tmp_path), str(src)],
                   check=True, capture_output=True, timeout=180)
    if not (tmp_path / "s.doc").exists():
        pytest.skip("LibreOffice Writer component not installed")
    text = load_document("s.doc", (tmp_path / "s.doc").read_bytes()).text
    assert "Bu sözleşme iki taraf arasında yapılmıştır." in text
